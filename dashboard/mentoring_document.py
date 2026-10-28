"""Build a filled Word mentoring record from the college's supplied template."""

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


TEMPLATE_PATH = Path(__file__).parent / "document_templates" / "mentorship_file_b.tech.docx"
SEMESTER_LABELS = [
    ("I B.Tech", "First Semester"),
    ("I B.Tech", "Second Semester"),
    ("II B.Tech", "First Semester"),
    ("II B.Tech", "Second Semester"),
    ("III B.Tech", "First Semester"),
    ("III B.Tech", "Second Semester"),
    ("IV B.Tech", "First Semester"),
    ("IV B.Tech", "Second Semester"),
]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%m-%Y")
    return str(value).strip()


def _set_cell_text(cell, value):
    """Replace cell text while retaining its existing first-run formatting."""
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0] if paragraphs else cell.add_paragraph()
    runs = paragraph.runs
    if runs:
        first_run = runs[0]
        first_run.text = _text(value)
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(_text(value))

    for extra_paragraph in paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)


def _selected_yes_no(value):
    return ("Yes [X]", "No [ ]") if value == "Yes" else ("Yes [ ]", "No [X]") if value == "No" else ("Yes [ ]", "No [ ]")


def _set_semester_radio(row, yes_cell_index, no_cell_index, value):
    yes_text, no_text = _selected_yes_no(value)
    _set_cell_text(row.cells[yes_cell_index], yes_text)
    _set_cell_text(row.cells[no_cell_index], no_text)


def _add_continuation_semester(document, semester_number, semester_data):
    year_label, semester_label = SEMESTER_LABELS[semester_number - 1]
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(4)
    heading_run = heading.add_run(f"{year_label} - {semester_label}")
    heading_run.bold = True
    heading_run.font.name = "Arial"
    heading_run.font.size = Pt(11)

    field_rows = [
        ("Academic Year", semester_data.get("academic_year")),
        ("Semester", semester_data.get("semester")),
        ("Name of the Mentor", semester_data.get("mentor_name")),
        ("Designation", semester_data.get("designation")),
        ("Semester Starting Date", semester_data.get("start_date")),
        ("Semester Ending Date", semester_data.get("end_date")),
        ("Hosteller / Day Scholar", semester_data.get("is_hosteller")),
        ("Travelling in College Bus", semester_data.get("college_bus")),
        ("Backlogs in previous Semester", semester_data.get("backlogs")),
        ("Backlog Courses", semester_data.get("backlog_courses")),
        ("Areas where support is required", semester_data.get("support_required")),
        ("Slot 1 - Observations by mentor", semester_data.get("slot1_observations")),
        ("Slot 1 - Suggestions given by mentor", semester_data.get("slot1_suggestions")),
        ("Slot 1 - Progress observed", semester_data.get("slot1_progress")),
        ("Slot 2 - Observations by mentor", semester_data.get("slot2_observations")),
        ("Slot 2 - Suggestions given by mentor", semester_data.get("slot2_suggestions")),
        ("Slot 2 - Progress observed", semester_data.get("slot2_progress")),
        ("Date of Mentoring", semester_data.get("date_of_mentoring")),
        ("Mentee Sign", semester_data.get("mentee_sign")),
        ("Mentor Sign", semester_data.get("mentor_sign")),
        ("Signature of AC", semester_data.get("signature_ac")),
        ("Signature of HOD", semester_data.get("signature_hod")),
    ]
    table = document.add_table(rows=len(field_rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, field_rows):
        row.cells[0].width = Inches(2.15)
        row.cells[1].width = Inches(4.0)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_text(row.cells[0], label)
        _set_cell_text(row.cells[1], value)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)


def create_mentoring_document(student, mentoring, photo_bytes=None):
    """Return a populated B.Tech mentoring record as DOCX bytes."""
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError("The B.Tech mentoring Word template is missing.")

    document = Document(TEMPLATE_PATH)
    tables = document.tables
    if len(tables) < 11:
        raise ValueError("The supplied mentoring template has an unexpected table layout.")

    semester_data = mentoring.semester_data if isinstance(mentoring.semester_data, dict) else {}
    student_name = _text(getattr(student, "student_name", ""))
    hall_ticket = _text(getattr(student, "ht_no", ""))
    branch = _text(getattr(mentoring, "branch", ""))
    batch = _text(getattr(mentoring, "batch", ""))
    admission_year = _text(getattr(mentoring, "admission_year", ""))

    # Student information sheet and identity details.
    _set_cell_text(tables[0].rows[0].cells[0], f"BATCH: {batch}")
    info = tables[1]
    _set_cell_text(info.rows[2].cells[2], batch)
    _set_cell_text(info.rows[2].cells[9], admission_year)
    identity = (
        f"Hall Ticket No.: {hall_ticket}    Name of the Student: {student_name}    "
        f"Program: B.Tech    Branch / Specialization: {branch}    "
        f"Re-admitted Student: {_text(mentoring.readmitted_student)}    "
        f"Lateral Entry Student: {_text(mentoring.lateral_entry_student)}    Photograph: Uploaded"
    )
    _set_cell_text(info.rows[3].cells[0], identity)
    if photo_bytes:
        photo_paragraph = info.rows[3].cells[0].add_paragraph()
        photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        photo_paragraph.add_run().add_picture(BytesIO(photo_bytes), width=Inches(1.0))
    _set_cell_text(info.rows[6].cells[5], mentoring.father_name)
    _set_cell_text(info.rows[8].cells[5], mentoring.mother_name)
    _set_cell_text(info.rows[11].cells[0], mentoring.address)
    _set_cell_text(info.rows[11].cells[3], mentoring.permanent_address)
    _set_cell_text(info.rows[11].cells[7], mentoring.communication_address)

    # Personal and contact details.
    _set_cell_text(tables[2].rows[1].cells[0], mentoring.gender)
    _set_cell_text(tables[2].rows[1].cells[1], mentoring.blood_group)
    _set_cell_text(tables[2].rows[1].cells[2], mentoring.allergies)
    _set_cell_text(tables[3].rows[1].cells[1], mentoring.student_mobile)
    _set_cell_text(tables[3].rows[1].cells[2], mentoring.student_email)
    _set_cell_text(tables[3].rows[2].cells[1], mentoring.parent_mobile)
    _set_cell_text(tables[3].rows[2].cells[2], mentoring.parent_email)

    # Educational background and grade summary.
    for row_index, prefix in ((1, "ssc"), (2, "inter")):
        fields = (
            f"{prefix}_school_name" if prefix == "ssc" else "inter_college_name",
            f"{prefix}_place" if prefix == "ssc" else "inter_place",
            f"{prefix}_board" if prefix == "ssc" else "inter_board",
            f"{prefix}_percentage" if prefix == "ssc" else "inter_percentage",
        )
        for cell_index, field_name in enumerate(fields, start=1):
            _set_cell_text(tables[4].rows[row_index].cells[cell_index], getattr(mentoring, field_name))

    for semester_number in range(1, 9):
        _set_cell_text(tables[5].rows[2].cells[semester_number], getattr(mentoring, f"sgpa_sem{semester_number}"))
        _set_cell_text(tables[5].rows[3].cells[semester_number], getattr(mentoring, f"cgpa_sem{semester_number}"))

    # Semester I details are built into the supplied template.
    sem1 = semester_data.get("1", {})
    sem1_table = tables[6]
    _set_cell_text(sem1_table.rows[1].cells[1], sem1.get("mentor_name"))
    _set_cell_text(sem1_table.rows[1].cells[8], sem1.get("designation"))
    _set_cell_text(
        sem1_table.rows[3].cells[0],
        f"From: {_text(sem1.get('start_date'))}    To: {_text(sem1.get('end_date'))}",
    )
    _set_semester_radio(sem1_table.rows[4], 2, 6, sem1.get("is_hosteller"))
    _set_semester_radio(sem1_table.rows[5], 2, 6, sem1.get("college_bus"))
    _set_cell_text(sem1_table.rows[9].cells[7], sem1.get("support_required"))
    for slot_number, row_index in ((1, 1), (2, 2)):
        slot = f"slot{slot_number}"
        for column_index, suffix in ((1, "observations"), (2, "suggestions"), (3, "progress")):
            _set_cell_text(tables[7].rows[row_index].cells[column_index], sem1.get(f"{slot}_{suffix}"))
    sem1_signoff = tables[8]
    for row_index, field_name in ((0, "date_of_mentoring"), (1, "mentee_sign"), (2, "mentor_sign")):
        _set_cell_text(sem1_signoff.rows[row_index].cells[1], sem1.get(field_name))
    _set_cell_text(
        sem1_signoff.rows[4].cells[0],
        f"Signature of AC: {_text(sem1.get('signature_ac'))}    Signature of HOD: {_text(sem1.get('signature_hod'))}",
    )

    # Semester II details are built into the supplied template.
    sem2 = semester_data.get("2", {})
    _set_cell_text(sem1_signoff.rows[6].cells[2], sem2.get("mentor_name"))
    _set_cell_text(sem1_signoff.rows[6].cells[8], sem2.get("designation"))
    _set_cell_text(sem1_signoff.rows[8].cells[0], sem2.get("start_date"))
    _set_cell_text(sem1_signoff.rows[8].cells[4], sem2.get("end_date"))
    _set_semester_radio(sem1_signoff.rows[9], 3, 7, sem2.get("is_hosteller"))
    _set_semester_radio(sem1_signoff.rows[10], 3, 7, sem2.get("college_bus"))
    _set_semester_radio(sem1_signoff.rows[11], 3, 7, sem2.get("backlogs"))
    _set_cell_text(sem1_signoff.rows[12].cells[3], sem2.get("backlog_courses"))
    _set_cell_text(sem1_signoff.rows[13].cells[3], sem2.get("support_required"))
    for slot_number, row_index in ((1, 1), (2, 2)):
        slot = f"slot{slot_number}"
        for column_index, suffix in ((1, "observations"), (2, "suggestions"), (3, "progress")):
            _set_cell_text(tables[9].rows[row_index].cells[column_index], sem2.get(f"{slot}_{suffix}"))
    for row_index, field_name in ((0, "date_of_mentoring"), (1, "mentee_sign"), (2, "mentor_sign")):
        _set_cell_text(tables[10].rows[row_index].cells[1], sem2.get(field_name))
    _set_cell_text(
        tables[10].rows[4].cells[0],
        f"Signature of AC: {_text(sem2.get('signature_ac'))}    Signature of HOD: {_text(sem2.get('signature_hod'))}",
    )

    # The first-semester part of the supplied template has no backlog fields.
    extra_heading = document.add_paragraph()
    extra_heading.paragraph_format.space_before = Pt(8)
    extra_heading.paragraph_format.space_after = Pt(3)
    extra_heading_run = extra_heading.add_run("I B.Tech - First Semester Additional Details")
    extra_heading_run.bold = True
    extra_heading_run.font.name = "Arial"
    extra_heading_run.font.size = Pt(10)
    extra_table = document.add_table(rows=2, cols=2)
    extra_table.style = "Table Grid"
    extra_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (label, value) in zip(
        extra_table.rows,
        (("Backlogs in previous Semester", sem1.get("backlogs")), ("Backlog Courses", sem1.get("backlog_courses"))),
    ):
        row.cells[0].width = Inches(2.15)
        row.cells[1].width = Inches(4.0)
        _set_cell_text(row.cells[0], label)
        _set_cell_text(row.cells[1], value)

    # The provided template contains only I B.Tech. semesters, so add the
    # remaining six records using matching, compact continuation tables.
    document.add_page_break()
    continuation_title = document.add_paragraph()
    continuation_title.paragraph_format.space_after = Pt(8)
    continuation_run = continuation_title.add_run("Remaining Semester Mentoring Records")
    continuation_run.bold = True
    continuation_run.font.name = "Arial"
    continuation_run.font.size = Pt(14)
    for semester_number in range(3, 9):
        _add_continuation_semester(
            document,
            semester_number,
            semester_data.get(str(semester_number), {}),
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
